# -*- coding: utf-8 -*-
"""
Bluebon 포인팅 오차 통합 분석
  PART 1 : 단일변수 상관분석 + 추세선 잔차 기반 이상치 시각화
            → diff_png/*.png  +  analysis_single_var.xlsx (단일변수_상관 탭 1개)
  PART 2 : 4가지 기준 이상치 분석
            → outlier_analysis/*.png  +  analysis_outlier_4criteria.xlsx
"""

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from scipy import stats
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import warnings
warnings.filterwarnings('ignore')
import matplotlib
matplotlib.rcParams['font.family'] = 'Malgun Gothic'
matplotlib.rcParams['axes.unicode_minus'] = False

# ============================================================
# 경로 설정
# ============================================================
EXCEL_PATH    = r"<DATA_ROOT>\10_BlueBON\pointing_diff\SatRev_metadata\1_merge\20260331\pointing_diff_analysis_20260331.xlsx"
DIFF_PNG_DIR  = r"<DATA_ROOT>\10_BlueBON\pointing_diff\SatRev_metadata\2_analysis\20260331\diff_png"
OUTLIER_DIR   = r"<DATA_ROOT>\10_BlueBON\pointing_diff\SatRev_metadata\2_analysis\20260331\outlier_analysis"
EXCEL_CORR    = r"<DATA_ROOT>\10_BlueBON\pointing_diff\SatRev_metadata\2_analysis\20260331\analysis_single_var.xlsx"
EXCEL_OUTLIER = r"<DATA_ROOT>\10_BlueBON\pointing_diff\SatRev_metadata\2_analysis\20260331\analysis_outlier_4criteria.xlsx"

Path(DIFF_PNG_DIR).mkdir(parents=True, exist_ok=True)
Path(OUTLIER_DIR).mkdir(parents=True, exist_ok=True)

# ============================================================
# 데이터 로드 (공통)
# ============================================================
df = pd.read_excel(EXCEL_PATH)
orig_cols = df.columns.tolist()
obs_time_col = 'obs_time.1' if 'obs_time.1' in df.columns else 'obs_time'
df = df.dropna(subset=[obs_time_col]).copy()
df['obs_dt'] = pd.to_datetime(df[obs_time_col], errors='coerce')
df = df.dropna(subset=['obs_dt']).reset_index(drop=True)
df['actual_lon'] = pd.to_numeric(df['actual_lon'], errors='coerce')

df['days_since_start'] = (df['obs_dt'] - df['obs_dt'].min()).dt.total_seconds() / 86400
df['utc_hour']         = df['obs_dt'].dt.hour + df['obs_dt'].dt.minute / 60
df['abs_along']        = df['along_diff(km)'].abs()
df['abs_across']       = df['across_diff(km)'].abs()
df['total_error_km']   = np.sqrt(df['along_diff(km)']**2 + df['across_diff(km)']**2)

telem = df.dropna(subset=['ECEF_Vx(m/s)']).copy() if 'ECEF_Vx(m/s)' in df.columns else df.iloc[0:0].copy()
n_all, n_telem = len(df), len(telem)

if len(telem):
    telem['Δalt']          = telem['Alt_GNSS(km)'] - telem['Alt_TLE(km)']
    telem['|Δalt|']        = telem['Δalt'].abs()
    telem['ang_rate_total']= np.sqrt(telem['AngRate_X(°/s)']**2 + telem['AngRate_Y(°/s)']**2 + telem['AngRate_Z(°/s)']**2)
    telem['ecef_speed']    = np.sqrt(telem['ECEF_Vx(m/s)']**2 + telem['ECEF_Vy(m/s)']**2 + telem['ECEF_Vz(m/s)']**2)
    telem['|ECEF_Vx|']    = telem['ECEF_Vx(m/s)'].abs()
    telem['|ECEF_Vy|']    = telem['ECEF_Vy(m/s)'].abs()
    telem['|ECEF_Vz|']    = telem['ECEF_Vz(m/s)'].abs()
    telem['|ECI_Vx|']     = telem['ECI_Vx(m/s)'].abs()
    telem['|ECI_Vy|']     = telem['ECI_Vy(m/s)'].abs()
    telem['|ECI_Vz|']     = telem['ECI_Vz(m/s)'].abs()
    telem['|AngRate_X|']  = telem['AngRate_X(°/s)'].abs()
    telem['|AngRate_Y|']  = telem['AngRate_Y(°/s)'].abs()
    telem['|AngRate_Z|']  = telem['AngRate_Z(°/s)'].abs()

    # 계산 컬럼을 df에도 반영
    for col in ['Δalt','|Δalt|','ang_rate_total','ecef_speed',
                '|ECEF_Vx|','|ECEF_Vy|','|ECEF_Vz|',
                '|ECI_Vx|','|ECI_Vy|','|ECI_Vz|',
                '|AngRate_X|','|AngRate_Y|','|AngRate_Z|']:
        if col not in df.columns:
            df[col] = np.nan
        df.loc[telem.index, col] = telem[col]

print(f"전체: {n_all}건, 텔레메트리: {n_telem}건")

def safe_filename(name):
    return (name.replace('/', '_per_').replace('\\', '_').replace('|', 'abs_')
                .replace('°', 'deg').replace('(', '_').replace(')', '_')
                .replace(' ', '').replace('Δ', 'delta_').replace('+', '_'))

# ============================================================
# PART 1 : 단일변수 상관분석 + 이상치 시각화
# ============================================================
print("\n" + "="*70)
print("PART 1: 단일변수 상관분석 + 추세선 잔차 기반 이상치 탐지")
print("="*70)

_all   = f'all({n_all})'
_telem = f'telem({n_telem})'

y_cols = {
    'along':   'along_diff(km)',
    'across':  'across_diff(km)',
    '|along|': 'abs_along',
    '|across|':'abs_across',
    'total':   'total_error_km',
}

x_all_list = [
    ('days_since_start', _all),
    ('target_lat',       _all),
    ('target_lon',       _all),
    ('utc_hour',         _all),
    ('Tilt (roll)',       _all),
    ('Pitch',            _all),
]
x_all_list = [(c, d) for c, d in x_all_list if c in df.columns]

x_telem_list = [
    ('Alt_GNSS(km)',      _telem), ('Alt_TLE(km)',      _telem),
    ('Δalt',              _telem), ('|Δalt|',           _telem),
    ('GPS_Sats',          _telem),
    ('AngRate_X(°/s)',    _telem), ('AngRate_Y(°/s)',   _telem), ('AngRate_Z(°/s)',   _telem),
    ('|AngRate_X|',       _telem), ('|AngRate_Y|',      _telem), ('|AngRate_Z|',      _telem),
    ('ang_rate_total',    _telem),
    ('ECEF_Vx(m/s)',      _telem), ('ECEF_Vy(m/s)',     _telem), ('ECEF_Vz(m/s)',     _telem),
    ('|ECEF_Vx|',         _telem), ('|ECEF_Vy|',        _telem), ('|ECEF_Vz|',        _telem),
    ('ecef_speed',        _telem),
    ('ECI_Vx(m/s)',       _telem), ('ECI_Vy(m/s)',      _telem), ('ECI_Vz(m/s)',      _telem),
    ('|ECI_Vx|',          _telem), ('|ECI_Vy|',         _telem), ('|ECI_Vz|',         _telem),
    ('TimeDiff_GNSS_OBC', _telem),
]
x_telem_list = [(c, d) for c, d in x_telem_list if c in telem.columns]

corr_results = []
graph_count  = 0

for y_name, y_col in y_cols.items():
    for x_col, dataset in x_all_list + x_telem_list:
        data  = telem if dataset == _telem else df
        valid = data[[x_col, y_col]].dropna()
        if len(valid) < 5 or np.std(valid[x_col].values) == 0 or np.std(valid[y_col].values) == 0:
            continue

        x, y = valid[x_col].values, valid[y_col].values
        rp, pp = stats.pearsonr(x, y)
        rs, ps = stats.spearmanr(x, y)
        rk, pk = stats.kendalltau(x, y)

        slope, intercept = np.polyfit(x, y, 1)
        residual  = y - (slope * x + intercept)
        resid_std = np.std(residual, ddof=1)
        is_outlier= np.abs(residual) > resid_std
        n_outlier = is_outlier.sum()

        corr_results.append({
            'Y': y_name, 'X': x_col, 'n': len(valid), 'dataset': dataset,
            'Pearson_r': rp, 'Pearson_p': pp,
            'Spearman_r': rs, 'Spearman_p': ps,
            'Kendall_tau': rk, 'Kendall_p': pk,
            'slope': slope, 'intercept': intercept,
            'resid_std': resid_std, 'n_outlier': n_outlier,
        })

        is_sig = min(pp, ps, pk) < 0.1
        fig, ax = plt.subplots(figsize=(8, 6))
        ax.scatter(x[~is_outlier], y[~is_outlier], c='#2166ac', alpha=0.7,
                   edgecolors='white', s=60, linewidths=0.5, label='정상', zorder=2)
        ax.scatter(x[is_outlier],  y[is_outlier],  c='#b2182b', alpha=0.9,
                   edgecolors='black', s=60, linewidths=0.8, marker='D',
                   label=f'이상치(|잔차|>1σ, {n_outlier}건)', zorder=3)
        x_line = np.linspace(x.min(), x.max(), 100)
        y_line = slope * x_line + intercept
        ax.plot(x_line, y_line, 'r-', linewidth=1.5, alpha=0.8,
                label=f'y = {slope:.4f}x + {intercept:.2f}')
        ax.fill_between(x_line, y_line - resid_std, y_line + resid_std,
                        alpha=0.12, color='red', label=f'±1σ ({resid_std:.2f}km)')
        ax.axhline(y=0, color='gray', linewidth=0.5)
        ax.text(0.03, 0.97,
                f'n={len(valid)}, outlier={n_outlier}\n'
                f'Pearson r={rp:.3f} (p={pp:.4f})\n'
                f'Spearman ρ={rs:.3f} (p={ps:.4f})\n'
                f'Kendall τ={rk:.3f} (p={pk:.4f})',
                transform=ax.transAxes, fontsize=10, verticalalignment='top',
                bbox=dict(boxstyle='round', facecolor='wheat', alpha=0.5))
        ax.set_xlabel(x_col, fontsize=12)
        ax.set_ylabel(y_col, fontsize=12)
        title_prefix = '★ ' if is_sig else ''
        t = ax.set_title(f'{title_prefix}{y_name}  vs  {x_col}  [{dataset}]',
                         fontsize=13, fontweight='bold')
        if is_sig:
            t.set_color('goldenrod')
        ax.legend(fontsize=9, loc='lower right')
        ax.grid(True, alpha=0.3)
        plt.tight_layout()
        prefix = 'SIG_' if is_sig else ''
        fname = f"{prefix}{safe_filename(y_name)}_vs_{safe_filename(x_col)}.png"
        fig.savefig(Path(DIFF_PNG_DIR) / fname, dpi=150, bbox_inches='tight')
        plt.close(fig)
        graph_count += 1
        sig = '***' if min(pp, ps, pk) < 0.05 else (' * ' if is_sig else '   ')
        print(f"  {sig} {fname} (outlier={n_outlier})")

df_corr = pd.DataFrame(corr_results)
print(f"\n총 {len(df_corr)}개 조합, 그래프 {graph_count}개")

print("\n[p < 0.05]")
for _, r in df_corr.iterrows():
    sigs = []
    if r['Pearson_p']  < 0.05: sigs.append(f"Pr={r['Pearson_r']:.3f}")
    if r['Spearman_p'] < 0.05: sigs.append(f"Sp={r['Spearman_r']:.3f}")
    if r['Kendall_p']  < 0.05: sigs.append(f"Kd={r['Kendall_tau']:.3f}")
    if sigs:
        print(f"  {r['Y']:>8s} vs {r['X']:>22s} [{r['dataset']}] {', '.join(sigs)} (outlier={r['n_outlier']})")

# ── PART 1 엑셀 저장 (단일변수_상관 탭 1개만) ──
print(f"\n엑셀 저장: {EXCEL_CORR}")
with pd.ExcelWriter(EXCEL_CORR, engine='openpyxl') as writer:
    df_corr.sort_values(['Y', 'Pearson_p']).to_excel(writer, sheet_name='단일변수_상관', index=False)

wb = load_workbook(EXCEL_CORR)
ws = wb['단일변수_상관']
hdr_fill = PatternFill('solid', fgColor='4472C4')
hdr_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
sig_fill  = PatternFill('solid', fgColor='C6EFCE')
bdr_fill  = PatternFill('solid', fgColor='FFEB9C')
dat_font  = Font(name='Arial', size=10)
thin      = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                   top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
hdr_vals  = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
p_cols    = [i+1 for i, v in enumerate(hdr_vals) if v and '_p' in str(v)]
for col in range(1, ws.max_column + 1):
    cell = ws.cell(row=1, column=col)
    cell.font, cell.fill = hdr_font, hdr_fill
    cell.alignment, cell.border = Alignment(horizontal='center'), thin
for row in range(2, ws.max_row + 1):
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        cell.font, cell.border = dat_font, thin
        cell.alignment = Alignment(horizontal='center')
        if isinstance(cell.value, float):
            cell.number_format = '0.0000'
    for pc in p_cols:
        cell = ws.cell(row=row, column=pc)
        if isinstance(cell.value, float):
            cell.fill = sig_fill if cell.value < 0.05 else (bdr_fill if cell.value < 0.1 else PatternFill())
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=ws.max_column).column_letter}{ws.max_row}"
wb.save(EXCEL_CORR)
print("  완료!")

# ============================================================
# PART 2 : 4가지 기준 이상치 분석
# ============================================================
print("\n" + "="*70)
print("PART 2: 4가지 기준 이상치 분석")
print("="*70)

params = [
    ('days_since_start',  '경과일'),
    ('target_lat',        '타겟위도'),
    ('target_lon',        '타겟경도'),
    ('utc_hour',          'UTC시간'),
    ('Tilt (roll)',        'Tilt(roll)'),
    ('Pitch',             'Pitch'),
    ('along_diff(km)',    'along오차'),
    ('across_diff(km)',   'across오차'),
    ('abs_along',         '|along|'),
    ('abs_across',        '|across|'),
    ('total_error_km',    'total오차'),
    ('Alt_GNSS(km)',      'GNSS고도'),
    ('Alt_TLE(km)',       'TLE고도'),
    ('Δalt',              'Δalt'),
    ('|Δalt|',            '|Δalt|'),
    ('GPS_Sats',          'GPS위성수'),
    ('AngRate_X(°/s)',    '각속도X'),
    ('AngRate_Y(°/s)',    '각속도Y'),
    ('AngRate_Z(°/s)',    '각속도Z'),
    ('|AngRate_X|',       '|각속도X|'),
    ('|AngRate_Y|',       '|각속도Y|'),
    ('|AngRate_Z|',       '|각속도Z|'),
    ('ang_rate_total',    '각속도크기'),
    ('ECEF_Vx(m/s)',      'ECEF_Vx'),
    ('ECEF_Vy(m/s)',      'ECEF_Vy'),
    ('ECEF_Vz(m/s)',      'ECEF_Vz'),
    ('|ECEF_Vx|',         '|ECEF_Vx|'),
    ('|ECEF_Vy|',         '|ECEF_Vy|'),
    ('|ECEF_Vz|',         '|ECEF_Vz|'),
    ('ecef_speed',        'ECEF속도크기'),
    ('TimeDiff_GNSS_OBC', 'GNSS-OBC시간차'),
]

# 기준 A/B: 추세선 잔차
slope_a, intercept_a = np.polyfit(df['days_since_start'], df['along_diff(km)'], 1)
df['along_pred']  = slope_a * df['days_since_start'] + intercept_a
df['along_resid'] = df['along_diff(km)'] - df['along_pred']
resid_std_a = df['along_resid'].std(ddof=1)

slope_b, intercept_b = np.polyfit(df['days_since_start'], df['across_diff(km)'], 1)
df['across_pred']  = slope_b * df['days_since_start'] + intercept_b
df['across_resid'] = df['across_diff(km)'] - df['across_pred']
resid_std_b = df['across_resid'].std(ddof=1)

# 기준 C/D: 절대값 임계
abs_along_thresh  = df['abs_along'].mean()  + df['abs_along'].std(ddof=1)
abs_across_thresh = df['abs_across'].mean() + df['abs_across'].std(ddof=1)

criteria = {
    'A_along추세선잔차': {
        'outlier_mask': df['along_resid'].abs() > resid_std_a,
        'desc': f'along vs days 추세선 |잔차|>{resid_std_a:.1f}km',
        'y_col': 'along_diff(km)', 'x_col': 'days_since_start',
        'slope': slope_a, 'intercept': intercept_a, 'resid_std': resid_std_a,
        'resid_col': 'along_resid', 'pred_col': 'along_pred',
    },
    'B_across추세선잔차': {
        'outlier_mask': df['across_resid'].abs() > resid_std_b,
        'desc': f'across vs days 추세선 |잔차|>{resid_std_b:.1f}km',
        'y_col': 'across_diff(km)', 'x_col': 'days_since_start',
        'slope': slope_b, 'intercept': intercept_b, 'resid_std': resid_std_b,
        'resid_col': 'across_resid', 'pred_col': 'across_pred',
    },
    'C_along절대값': {
        'outlier_mask': df['abs_along'] > abs_along_thresh,
        'desc': f'|along|>{abs_along_thresh:.1f}km (mean+1σ)',
        'y_col': 'along_diff(km)', 'x_col': None, 'threshold': abs_along_thresh,
    },
    'D_across절대값': {
        'outlier_mask': df['abs_across'] > abs_across_thresh,
        'desc': f'|across|>{abs_across_thresh:.1f}km (mean+1σ)',
        'y_col': 'across_diff(km)', 'x_col': None, 'threshold': abs_across_thresh,
    },
}

all_range_rows    = []
all_outlier_sheets= {}

for crit_name, crit in criteria.items():
    mask     = crit['outlier_mask']
    outliers = df[mask].copy()
    normals  = df[~mask].copy()

    print(f"\n{'='*70}")
    print(f"  {crit_name}: {crit['desc']}")
    print(f"  이상치 {len(outliers)}건, 정상 {len(normals)}건")
    print(f"{'='*70}")

    # 이상치 목록 출력
    print("\n  [이상치 목록]")
    for _, r in outliers.iterrows():
        resid_str = (f"잔차={r[crit['resid_col']]:>+7.1f}km" if 'resid_col' in crit
                     else f"|값|={r['abs_along' if 'along' in crit_name else 'abs_across']:>6.1f}km")
        tl = (f"GNSS={r['Alt_GNSS(km)']:.0f} GPS={r['GPS_Sats']:.0f} Vz={r['ECEF_Vz(m/s)']:.0f}"
              if pd.notna(r.get('Alt_GNSS(km)')) else "텔레메트리없음")
        print(f"    {r['obs_dt'].strftime('%m/%d %H:%M')} | {str(r['site']):30s} | "
              f"along={r['along_diff(km)']:>7.1f} across={r['across_diff(km)']:>6.1f} | {resid_str} | {tl}")

    # 이상치 목록 시트용 컬럼 구성
    base_list_cols = ['site', 'obs_dt', 'along_diff(km)', 'across_diff(km)', 'total_error_km',
                      'target_lat', 'target_lon', 'Tilt (roll)', 'Pitch',
                      'days_since_start', 'utc_hour']
    if 'resid_col' in crit:
        outliers['잔차(km)']  = outliers[crit['resid_col']]
        outliers['예측값(km)'] = outliers[crit['pred_col']]
        base_list_cols = ['site', 'obs_dt', crit['y_col'], '예측값(km)', '잔차(km)',
                          'along_diff(km)', 'across_diff(km)', 'total_error_km',
                          'target_lat', 'target_lon', 'Tilt (roll)', 'Pitch',
                          'days_since_start', 'utc_hour']
    telem_cols_list = ['Alt_GNSS(km)', 'Alt_TLE(km)', 'Δalt', '|Δalt|', 'GPS_Sats',
                       'AngRate_X(°/s)', 'AngRate_Y(°/s)', 'AngRate_Z(°/s)', 'ang_rate_total',
                       'ECEF_Vx(m/s)', 'ECEF_Vy(m/s)', 'ECEF_Vz(m/s)',
                       '|ECEF_Vx|', '|ECEF_Vy|', '|ECEF_Vz|', 'ecef_speed',
                       'TimeDiff_GNSS_OBC']
    list_cols = base_list_cols + [c for c in telem_cols_list if c in outliers.columns]
    list_cols = [c for c in list_cols if c in outliers.columns]
    seen, unique_cols = set(), []
    for c in list_cols:
        if c not in seen:
            seen.add(c); unique_cols.append(c)

    sort_col = crit.get('resid_col', 'abs_along' if 'along' in crit_name else 'abs_across')
    all_outlier_sheets[crit_name] = outliers.sort_values(sort_col, key=abs, ascending=False)[unique_cols]

    # 파라미터 범위 비교 (t-test)
    print(f"\n  [파라미터 범위 비교]")
    print(f"  {'파라미터':>16s} | {'이상치 min~max (mean)':>30s} | {'정상 min~max (mean)':>30s} | {'p':>8s} | 해석")
    print(f"  {'-'*110}")
    for p_col, p_name in params:
        if p_col not in df.columns: continue
        o = outliers[p_col].dropna(); n_s = normals[p_col].dropna()
        if len(o) < 1 or len(n_s) < 1: continue
        tp = stats.ttest_ind(o, n_s, equal_var=False)[1] if len(o) >= 2 and len(n_s) >= 2 else np.nan
        interp = ('★★ 유의' if pd.notna(tp) and tp < 0.05 else
                  '★ 경계'  if pd.notna(tp) and tp < 0.1  else '')
        if interp:
            print(f"  {p_name:>16s} | {o.min():>9.1f}~{o.max():>8.1f} ({o.mean():>8.1f}) | "
                  f"{n_s.min():>9.1f}~{n_s.max():>8.1f} ({n_s.mean():>8.1f}) | "
                  f"{'n/a':>8s} | {interp}" if not pd.notna(tp) else
                  f"  {p_name:>16s} | {o.min():>9.1f}~{o.max():>8.1f} ({o.mean():>8.1f}) | "
                  f"{n_s.min():>9.1f}~{n_s.max():>8.1f} ({n_s.mean():>8.1f}) | "
                  f"{tp:>8.4f} | {interp}")
        all_range_rows.append({
            '기준': crit_name, '기준설명': crit['desc'],
            '파라미터': p_name, '파라미터_컬럼': p_col,
            '이상치_n': len(o), '이상치_min': o.min(), '이상치_max': o.max(),
            '이상치_mean': o.mean(), '이상치_std': o.std() if len(o) >= 2 else np.nan,
            '정상_n': len(n_s), '정상_min': n_s.min(), '정상_max': n_s.max(),
            '정상_mean': n_s.mean(), '정상_std': n_s.std() if len(n_s) >= 2 else np.nan,
            'p_value': tp,
        })

    # 그래프 1: 추세선 (기준 A, B)
    if 'slope' in crit:
        fig, ax = plt.subplots(figsize=(10, 7))
        ax.scatter(normals['days_since_start'], normals[crit['y_col']],
                   c='#2166ac', alpha=0.7, edgecolors='white', s=60, linewidths=0.5,
                   label=f'정상 (n={len(normals)})', zorder=2)
        ax.scatter(outliers['days_since_start'], outliers[crit['y_col']],
                   c='#b2182b', alpha=0.9, edgecolors='black', s=80, linewidths=0.8,
                   marker='D', label=f'이상치 (n={len(outliers)})', zorder=3)
        x_line = np.linspace(df['days_since_start'].min(), df['days_since_start'].max(), 100)
        y_line = crit['slope'] * x_line + crit['intercept']
        ax.plot(x_line, y_line, 'r-', linewidth=2, alpha=0.8,
                label=f'y={crit["slope"]:.4f}x+{crit["intercept"]:.2f}')
        ax.fill_between(x_line, y_line - crit['resid_std'], y_line + crit['resid_std'],
                        alpha=0.12, color='red', label=f'±1σ ({crit["resid_std"]:.2f}km)')
        ax.axhline(y=0, color='gray', linewidth=0.5)
        for _, r in outliers.iterrows():
            ax.annotate(str(r['site'])[:18], (r['days_since_start'], r[crit['y_col']]),
                        fontsize=7, alpha=0.8, xytext=(5, 5), textcoords='offset points')
        ax.set_xlabel('days_since_start', fontsize=12)
        ax.set_ylabel(crit['y_col'], fontsize=12)
        ax.set_title(f'{crit_name}: {crit["desc"]}', fontsize=13, fontweight='bold')
        ax.legend(fontsize=9, loc='lower right'); ax.grid(True, alpha=0.3)
        plt.tight_layout()
        fname = f'trendline_{safe_filename(crit_name)}.png'
        fig.savefig(Path(OUTLIER_DIR) / fname, dpi=150, bbox_inches='tight')
        plt.close(fig)
        print(f"\n  그래프: {fname}")

    # 그래프 2: 절대값 임계 (기준 C, D)
    if 'threshold' in crit:
        err_col = 'abs_along' if 'along' in crit_name else 'abs_across'
        fig, ax = plt.subplots(figsize=(10, 7))
        ax.scatter(normals['days_since_start'], normals[err_col],
                   c='#2166ac', alpha=0.7, edgecolors='white', s=60, linewidths=0.5,
                   label=f'정상 (n={len(normals)})', zorder=2)
        ax.scatter(outliers['days_since_start'], outliers[err_col],
                   c='#b2182b', alpha=0.9, edgecolors='black', s=80, linewidths=0.8,
                   marker='D', label=f'이상치 (n={len(outliers)})', zorder=3)
        ax.axhline(y=crit['threshold'], color='red', linewidth=1.5, linestyle='--',
                   label=f'mean+1σ = {crit["threshold"]:.1f}km')
        for _, r in outliers.iterrows():
            ax.annotate(str(r['site'])[:18], (r['days_since_start'], r[err_col]),
                        fontsize=7, alpha=0.8, xytext=(5, 5), textcoords='offset points')
        ax.set_xlabel('days_since_start', fontsize=12)
        ax.set_ylabel(err_col, fontsize=12)
        ax.set_title(f'{crit_name}: {crit["desc"]}', fontsize=13, fontweight='bold')
        ax.legend(fontsize=9); ax.grid(True, alpha=0.3)
        plt.tight_layout()
        fname = f'threshold_{safe_filename(crit_name)}.png'
        fig.savefig(Path(OUTLIER_DIR) / fname, dpi=150, bbox_inches='tight')
        plt.close(fig)
        print(f"  그래프: {fname}")

    # 그래프 3: 박스플롯
    telem_key = 'ECEF_Vx(m/s)' if 'ECEF_Vx(m/s)' in df.columns else None
    o_t = outliers.dropna(subset=[telem_key]) if telem_key else outliers.iloc[0:0]
    n_t = normals.dropna(subset=[telem_key])  if telem_key else normals.iloc[0:0]

    plot_params = [(c, n) for c, n in [
        ('Alt_GNSS(km)', 'GNSS고도'), ('|Δalt|', '|Δalt|'), ('GPS_Sats', 'GPS위성수'),
        ('ang_rate_total', '각속도크기'), ('ECEF_Vz(m/s)', 'ECEF_Vz'),
        ('|ECEF_Vz|', '|ECEF_Vz|'), ('ECEF_Vy(m/s)', 'ECEF_Vy'), ('ecef_speed', 'ECEF속도'),
        ('Tilt (roll)', 'Tilt(roll)'), ('Pitch', 'Pitch'), ('TimeDiff_GNSS_OBC', 'GNSS-OBC시간차'),
    ] if c in df.columns]

    if len(o_t) >= 1 and len(n_t) >= 1:
        ncols = 4
        nrows = (len(plot_params) + ncols - 1) // ncols
        fig, axes = plt.subplots(nrows, ncols, figsize=(ncols * 3.5, nrows * 4))
        axes = axes.flatten()
        for i, (p_col, p_name) in enumerate(plot_params):
            ax = axes[i]
            o = o_t[p_col].dropna() if p_col in o_t.columns else pd.Series(dtype=float)
            n = n_t[p_col].dropna() if p_col in n_t.columns else pd.Series(dtype=float)
            if len(o) < 1 or len(n) < 1:
                ax.set_visible(False); continue
            bp = ax.boxplot([n.values, o.values], labels=['정상', '이상치'],
                            patch_artist=True, widths=0.5)
            bp['boxes'][0].set_facecolor('#C6EFCE')
            bp['boxes'][1].set_facecolor('#FFC7CE')
            ax.scatter(np.ones(len(n)) + np.random.normal(0, 0.04, len(n)),
                       n.values, c='#2166ac', alpha=0.5, s=25, zorder=3)
            o_jitter = np.random.normal(0, 0.04, len(o))
            o_x = np.ones(len(o)) * 2 + o_jitter
            ax.scatter(o_x, o.values, c='#b2182b', alpha=0.7, s=40, zorder=3)
            # 이상치 날짜/시간 레이블
            if 'obs_dt' in o_t.columns:
                o_dt = o_t.loc[o.index, 'obs_dt']
                for xi, yi, dt in zip(o_x, o.values, o_dt):
                    if pd.notna(dt):
                        ax.annotate(pd.Timestamp(dt).strftime('%m/%d\n%H:%M'),
                                    xy=(xi, yi), fontsize=5.5, alpha=0.85,
                                    ha='left', va='bottom',
                                    xytext=(3, 2), textcoords='offset points')
            sig = (f'p={stats.ttest_ind(o, n, equal_var=False)[1]:.3f}' +
                   (' ★★' if stats.ttest_ind(o, n, equal_var=False)[1] < 0.05 else
                    ' ★'  if stats.ttest_ind(o, n, equal_var=False)[1] < 0.1  else '')
                   if len(o) >= 2 and len(n) >= 2 else f'n={len(o)}')
            ax.set_title(f'{p_name}\n{sig}', fontsize=9, fontweight='bold')
            ax.grid(True, alpha=0.3, axis='y')
        for j in range(len(plot_params), len(axes)):
            axes[j].set_visible(False)
        fig.suptitle(f'{crit_name}: 이상치(텔레{len(o_t)}건) vs 정상({len(n_t)}건)',
                     fontsize=12, fontweight='bold')
        plt.tight_layout()
        fname = f'boxplot_{safe_filename(crit_name)}.png'
        fig.savefig(Path(OUTLIER_DIR) / fname, dpi=150, bbox_inches='tight')
        plt.close(fig)
        print(f"  그래프: {fname}")

# ── PART 2 엑셀 저장 ──
print(f"\n{'='*70}")
print(f"엑셀 저장: {EXCEL_OUTLIER}")
df_range = pd.DataFrame(all_range_rows)
with pd.ExcelWriter(EXCEL_OUTLIER, engine='openpyxl') as writer:
    df_range.to_excel(writer, sheet_name='파라미터_범위비교', index=False)
    for crit_name, df_list in all_outlier_sheets.items():
        df_list.to_excel(writer, sheet_name=crit_name[:31], index=False)

wb2 = load_workbook(EXCEL_OUTLIER)
hdr_fill2 = PatternFill('solid', fgColor='4472C4')
hdr_font2 = Font(name='Arial', bold=True, color='FFFFFF', size=10)
sig_fill2  = PatternFill('solid', fgColor='C6EFCE')
bdr_fill2  = PatternFill('solid', fgColor='FFEB9C')
dat_font2  = Font(name='Arial', size=10)
thin2      = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
for ws_name in wb2.sheetnames:
    ws2 = wb2[ws_name]
    hdr_vals2 = [ws2.cell(row=1, column=c).value for c in range(1, ws2.max_column + 1)]
    p_idx2    = [i+1 for i, v in enumerate(hdr_vals2) if v and 'p_value' in str(v)]
    for col in range(1, ws2.max_column + 1):
        cell = ws2.cell(row=1, column=col)
        cell.font, cell.fill = hdr_font2, hdr_fill2
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = thin2
    for row in range(2, ws2.max_row + 1):
        for col in range(1, ws2.max_column + 1):
            cell = ws2.cell(row=row, column=col)
            cell.font, cell.border = dat_font2, thin2
            cell.alignment = Alignment(horizontal='center')
            if isinstance(cell.value, float):
                cell.number_format = '0.00'
        for pc in p_idx2:
            cell = ws2.cell(row=row, column=pc)
            if isinstance(cell.value, float):
                cell.fill = sig_fill2 if cell.value < 0.05 else (bdr_fill2 if cell.value < 0.1 else PatternFill())
    ws2.freeze_panes = 'A2'
    ws2.auto_filter.ref = f"A1:{ws2.cell(row=1, column=ws2.max_column).column_letter}{ws2.max_row}"
    for col in range(1, ws2.max_column + 1):
        ws2.column_dimensions[ws2.cell(row=1, column=col).column_letter].width = 14
wb2.save(EXCEL_OUTLIER)

print("\n완료!")
print(f"  PART 1 그래프 → {DIFF_PNG_DIR}")
print(f"  PART 1 엑셀   → {EXCEL_CORR}  (단일변수_상관 탭)")
print(f"  PART 2 그래프 → {OUTLIER_DIR}")
print(f"  PART 2 엑셀   → {EXCEL_OUTLIER}  (파라미터_범위비교 + 4기준 탭)")
